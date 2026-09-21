import openpyxl
import os
import re
import shutil
import time

EXCEL_PATH = r'd:\Kerja praktek\dse automation\tracker-posting.xlsx'

def load_blacklist():
    blacklist_path = os.path.join(os.path.dirname(__file__), 'blacklist.txt')
    if not os.path.exists(blacklist_path):
        return set()
    try:
        with open(blacklist_path, 'r') as f:
            return set(line.strip().lower() for line in f if line.strip())
    except Exception:
        return set()

def categorize_site(domain, blacklist=None):
    domain_lower = domain.lower() if domain else ''
    
    if blacklist is not None and domain_lower in blacklist:
        return "Blacklisted"
    
    image_keywords = ['img', 'pic', 'image', 'photo', 'ibb', 'postimg', 'prnt', 'shutterfly', 'gifyu', 'imgur', 'freeimage', 'catbox', 'pixeldrain', 'uguu', 'gofile', 'photobucket', 'yogile', 'flickr', 'lensdump', 'pixhost', 'chevereto', 'smugmug', 'sirv', 'bytescale', 'transloadit', 'uploadthing', 'snipboard', 'turboimagehost', 'pasteboard', 'myalbum']
    paste_keywords = ['paste', 'telegra.ph', 'justpaste', 'jpst', 'hubpages', 'write.as', 'writeupcafe', 'notion', 'padlet']
    cloud_keywords = ['cloud', 'pcloud', 'mediafire', '4shared', 'mega.nz', 'drive', 'dropbox', 'cloudup', 'koofr', 'pinata', 'storacha', 'filebase', 'backblaze', 'storj', 'wasabi', 'supabase', 'appwrite', 'firebase', 'onedrive', 'box', 'proton', 'internxt', 'icedrive', 'filen', 'sync', 'idrive', 'terabox', 'degoo', 'blomp', 'opendrive', 'yandex', 'jottacloud', 'nordlocker', 'cozy', 'mobidrive', 'digiboxx', 'icloud']
    blog_keywords = ['medium', 'blogger', 'blogspot', 'tumblr', 'livejournal', 'wattpad', 'wix', 'penzu', 'weebly', 'ghost', 'squarespace', 'wordpress', 'strikingly', 'site123', 'jimdo', 'webnode', 'ucraft', 'dorik', 'typedream', 'tilda', 'framer', 'substack', 'hubspot', 'mozello', 'yola', 'webador', 'simdif', 'imcreator', 'journoportfolio', 'pixpa']
    
    if any(k in domain_lower for k in image_keywords):
        return "Image Host"
    elif any(k in domain_lower for k in paste_keywords):
        return "Text / Paste Host"
    elif any(k in domain_lower for k in cloud_keywords):
        return "Cloud Storage"
    elif any(k in domain_lower for k in blog_keywords):
        return "Blog / CMS"
    else:
        return "Social / Web Platform"

class TrackerManager:
    def __init__(self, excel_path=EXCEL_PATH):
        self.excel_path = excel_path

    def _safe_load_workbook(self, data_only=True):
        if not os.path.exists(self.excel_path):
            return None
        
        try:
            return openpyxl.load_workbook(self.excel_path, data_only=data_only)
        except Exception:
            # If Excel file is opened and locked by MS Excel, copy to temp file and read
            temp_copy = self.excel_path + ".read.tmp"
            try:
                shutil.copyfile(self.excel_path, temp_copy)
                wb = openpyxl.load_workbook(temp_copy, data_only=data_only)
                try:
                    os.remove(temp_copy)
                except Exception:
                    pass
                return wb
            except Exception as e:
                print(f"Error opening excel file safely: {e}")
                return None

    def get_data(self):
        wb = self._safe_load_workbook(data_only=True)
        if not wb:
            return {"headers": [], "rows": [], "campaigns": [], "stats": {}}
        
        sheet = wb['Tracker'] if 'Tracker' in wb.sheetnames else wb.active
        
        headers = [sheet.cell(row=1, column=j).value for j in range(1, sheet.max_column + 1)]
        headers = [str(h).strip() for h in headers if h is not None]
        
        campaigns = headers[1:] if len(headers) > 1 else []
        
        rows = []
        category_counts = {
            "Image Host": 0,
            "Text / Paste Host": 0,
            "Cloud Storage": 0,
            "Blog / CMS": 0,
            "Social / Web Platform": 0,
            "Blacklisted": 0
        }
        
        total_posts_per_campaign = {c: 0 for c in campaigns}
        
        blacklist = load_blacklist()
        
        for i in range(2, sheet.max_row + 1):
            situs = sheet.cell(row=i, column=1).value
            if not situs:
                continue
            
            situs = str(situs).strip()
            category = categorize_site(situs, blacklist=blacklist)
            category_counts[category] = category_counts.get(category, 0) + 1
            
            campaign_links = {}
            for idx, c in enumerate(campaigns):
                col_idx = idx + 2
                val = sheet.cell(row=i, column=col_idx).value
                link = str(val).strip() if val else None
                campaign_links[c] = link
                if link:
                    total_posts_per_campaign[c] += 1
            
            rows.append({
                "row_id": i,
                "situs": situs,
                "category": category,
                "campaigns": campaign_links
            })
            
        stats = {
            "total_sites": len(rows),
            "category_counts": category_counts,
            "total_posts_per_campaign": total_posts_per_campaign,
            "total_campaigns": len(campaigns)
        }
        
        return {
            "headers": headers,
            "campaigns": campaigns,
            "rows": rows,
            "stats": stats
        }

    def add_campaign(self, campaign_name):
        wb = self._safe_load_workbook(data_only=False)
        if not wb:
            return False, "Failed to access Excel file"

        sheet = wb['Tracker'] if 'Tracker' in wb.sheetnames else wb.active
        
        headers = [sheet.cell(row=1, column=j).value for j in range(1, sheet.max_column + 1)]
        headers = [str(h).strip() for h in headers if h is not None]

        if campaign_name in headers:
            return False, f"Campaign '{campaign_name}' already exists"
        
        new_col = len(headers) + 1
        sheet.cell(row=1, column=new_col).value = campaign_name
        
        try:
            wb.save(self.excel_path)
            return True, f"Campaign '{campaign_name}' added successfully"
        except PermissionError:
            return False, "Permission denied: Please close tracker-posting.xlsx in Microsoft Excel before saving!"
        except Exception as e:
            return False, f"Save failed: {str(e)}"

    def update_cell(self, row_id, campaign_name, link):
        wb = self._safe_load_workbook(data_only=False)
        if not wb:
            return False, "Failed to access Excel file"

        sheet = wb['Tracker'] if 'Tracker' in wb.sheetnames else wb.active
        
        headers = [sheet.cell(row=1, column=j).value for j in range(1, sheet.max_column + 1)]
        headers = [str(h).strip() for h in headers if h is not None]

        if campaign_name not in headers:
            return False, "Campaign not found"
        
        col_idx = headers.index(campaign_name) + 1
        sheet.cell(row=int(row_id), column=col_idx).value = link

        try:
            wb.save(self.excel_path)
            return True, "Updated successfully"
        except PermissionError:
            return False, "Permission denied: Please close tracker-posting.xlsx in Microsoft Excel before saving!"
        except Exception as e:
            return False, f"Save failed: {str(e)}"

    def update_batch_results(self, campaign_name, results_dict):
        """
        results_dict: {situs_domain: link_url}
        """
        wb = self._safe_load_workbook(data_only=False)
        if not wb:
            return 0

        sheet = wb['Tracker'] if 'Tracker' in wb.sheetnames else wb.active
        
        headers = [sheet.cell(row=1, column=j).value for j in range(1, sheet.max_column + 1)]
        headers = [str(h).strip() for h in headers if h is not None]

        if campaign_name not in headers:
            headers.append(campaign_name)
            col_idx = len(headers)
            sheet.cell(row=1, column=col_idx).value = campaign_name
        else:
            col_idx = headers.index(campaign_name) + 1
            
        updated_count = 0
        for i in range(2, sheet.max_row + 1):
            situs = sheet.cell(row=i, column=1).value
            if not situs:
                continue
            situs_str = str(situs).strip().lower()
            
            for key_domain, link_url in results_dict.items():
                key_clean = key_domain.strip().lower()
                if key_clean in situs_str or situs_str in key_clean:
                    sheet.cell(row=i, column=col_idx).value = link_url
                    updated_count += 1
                    break
                    
        try:
            wb.save(self.excel_path)
        except Exception as e:
            print(f"Error saving batch results to excel: {e}")

        return updated_count
